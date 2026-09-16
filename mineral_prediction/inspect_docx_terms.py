from pathlib import Path
import re
import sys

from docx import Document


def main():
    path = Path(sys.argv[1])
    terms = [
        "multi-ring",
        "multiple ring",
        "ring buffer",
        "annulus",
        "buffer",
        "fault",
        "density",
        "distance",
        "grid",
    ]
    document = Document(path)
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    for index, text in enumerate(blocks):
        normalized = re.sub(r"\s+", " ", text).strip()
        if not normalized:
            continue
        matches = [term for term in terms if term.lower() in normalized.lower()]
        if matches:
            print(f"{index}: terms={matches}: {normalized}")


if __name__ == "__main__":
    main()
