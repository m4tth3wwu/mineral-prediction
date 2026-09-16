from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_method_comparison_report import (
    BLUE,
    DARK_BLUE,
    LIGHT_BLUE,
    MID_GRAY,
    add_body,
    add_callout,
    add_heading,
    add_table,
    load_results,
    mean_sd,
    set_run_font,
)


WORKSPACE = Path(__file__).resolve().parent
OUT_DIR = WORKSPACE / "output" / "method_comparison_report"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "传统统计与ACAWLR结果简单说明.docx"


def build():
    cv, external, deposit, point_metrics, coefficients, _, diagnostics = load_results()
    summary = mean_sd(cv)
    ext = external.set_index("model")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color, before, after in (
        ("Heading 1", 15, BLUE, 12, 6),
        ("Heading 2", 12, DARK_BLUE, 8, 4),
        ("Heading 3", 11, DARK_BLUE, 6, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("传统统计与ACAWLR结果简单说明")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("原论文44个训练矿床点与16个外部验证矿床点")
    set_run_font(r, size=11, color=MID_GRAY)

    add_callout(
        doc,
        "一句话结论",
        "ACAWLR的主要优势是排序能力更好、外部误报更少、预测为矿床时更可靠；传统Logistic更简单、更容易解释，并且在当前阈值下召回率和验证矿床命中数略高。两者外部F1几乎相同。",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "1. 数据怎么处理", 1)
    add_body(doc, "因变量是矿床是否存在（deposit/Label），铜Cu是自变量。研究区划分为1 km网格，44个训练矿床周围2 km内的网格作为训练正样本；16个验证矿床只用于最后的外部验证，不参加训练、标准化或阈值选择。")
    add_table(
        doc,
        ["数据", "矿床数", "正网格", "背景网格", "总网格"],
        [["训练", 44, 572, "11,440", "12,012"], ["外部验证", 16, 208, "4,160", "4,368"]],
        [2100, 1200, 1700, 2100, 2260],
        numeric_cols={1, 2, 3, 4},
        font_size=9,
    )
    add_body(doc, "背景网格是研究区内抽取的对照位置，并不是经过勘查确认的“绝对无矿点”。所以模型输出更适合解释为相对找矿有利度，而不是自然条件下的真实矿床发生概率。")

    add_heading(doc, "2. 十二邻点IDW是什么", 1)
    add_body(doc, "对每个1 km网格中心，从沉积物地球化学数据库中寻找距离最近的12个采样点。这12个点完全按投影后的空间距离选择，不是从44个矿床点中选择，也不根据铜含量高低选择。")
    add_body(doc, "插值权重为w=1/max(d,1 km)²，距离越近权重越大。Cu、Au、Mo、Fe使用同一组12个空间近邻分别插值；某元素缺失时只使用该元素的有效邻点并重新归一化。如果12点中该元素全部缺失，后续用训练集该变量的中位数填补。当前代码没有最大搜索半径，因此采样稀疏处的第12个点可能较远。")

    add_heading(doc, "3. 十个自变量", 1)
    rows = [
        ["Cu_ppm", "铜含量，反映铜地球化学异常"],
        ["Au_ppm", "金含量，反映金及伴生矿化异常"],
        ["Mo_ppm", "钼含量，反映钼及斑岩/热液矿化异常"],
        ["Fe_pct", "铁含量，反映蚀变和地质背景"],
        ["dist_fault", "到最近断层的距离，越小表示越靠近断层"],
        ["fault_density", "50 km范围内断层总长度/面积，表示构造发育程度"],
        ["local_fault_confidence", "断层方向集中度0–1，越大表示附近断层方向越一致"],
        ["gravity_bouguer", "Bouguer重力异常，反映地下密度差异"],
        ["gravity_distance_deg", "到最近重力点的角距离，是数据覆盖质量指标"],
        ["has_gravity", "重力值是否在0.25°内有效匹配，取0或1"],
    ]
    add_table(doc, ["变量", "含义"], rows, [2500, 6860], font_size=8.8)
    add_body(doc, "主网格分析中，十个变量都进行训练集内中位数填补和Z标准化，但没有统一取对数。标准化公式是z=(x−训练均值)/训练标准差。另有local_fault_angle只用于ACAWLR构造断层方向性空间距离，不属于十个回归变量。")

    add_heading(doc, "4. 两种方法有什么区别", 1)
    add_heading(doc, "传统全局Logistic", 2)
    add_body(doc, "模型假定整个研究区使用同一组固定系数：logit(p)=β₀+Σβⱼzⱼ。采用L2正则化和类别平衡权重，优点是结构简单、计算快、变量方向容易解释；不足是无法表示同一地质变量在不同空间位置作用强度不同。")
    add_heading(doc, "ACAWLR", 2)
    add_body(doc, "ACAWLR先用全局Logistic系数作为基础，再结合样本位置和局部断层方向生成空间变化的系数。可以概括为βᵢ=wᵢ⊙β_global，即每个位置都有一组经过空间网络调整的局部系数，因此能够学习空间非平稳性和断层方向性。")

    add_heading(doc, "5. 五折、排序和阈值", 1)
    add_body(doc, "五折交叉验证不是把网格随意分成五份：同一矿床缓冲区内的正网格必须在同一折，背景网格按100 km空间块分组，从而减少相邻网格跨折造成的数据泄漏。")
    add_body(doc, "AUC和PR-AUC评价概率排序，不需要阈值。Precision、Recall和F1要先选择阈值，把概率分成0或1。最终外部验证阈值来自44点训练集的五折OOF预测中使F1最大的值：ACAWLR为0.867，Logistic为0.918。")

    add_heading(doc, "6. 结果比较", 1)
    cv_rows = []
    for model in ["ACAWLR", "Global logistic"]:
        cv_rows.append([
            model,
            f"{summary.loc[model, ('auc', 'mean')]:.3f}±{summary.loc[model, ('auc', 'std')]:.3f}",
            f"{summary.loc[model, ('pr_auc', 'mean')]:.3f}±{summary.loc[model, ('pr_auc', 'std')]:.3f}",
            f"{summary.loc[model, ('precision', 'mean')]:.3f}",
            f"{summary.loc[model, ('recall', 'mean')]:.3f}",
            f"{summary.loc[model, ('f1', 'mean')]:.3f}",
        ])
    add_table(doc, ["五折结果", "AUC", "PR-AUC", "Precision", "Recall", "F1"], cv_rows,
              [1900, 1500, 1600, 1500, 1400, 1460], numeric_cols={1, 2, 3, 4, 5}, font_size=8.8)

    ext_rows = []
    for model in ["ACAWLR", "Global logistic"]:
        r = ext.loc[model]
        ext_rows.append([
            model, f"{r['auc']:.3f}", f"{r['pr_auc']:.3f}", f"{r['precision']:.3f}",
            f"{r['recall']:.3f}", f"{r['f1']:.3f}", int(r["fp"]), int(r["tp"]),
        ])
    add_table(doc, ["外部验证", "AUC", "PR-AUC", "Precision", "Recall", "F1", "FP", "TP"], ext_rows,
              [1800, 1050, 1200, 1400, 1200, 1100, 800, 810], numeric_cols={1, 2, 3, 4, 5, 6, 7}, font_size=8.6)
    ac_hits = int(deposit["acawlr_detected"].sum())
    log_hits = int(deposit["global_logistic_detected"].sum())
    add_body(doc, f"五折中，ACAWLR的平均AUC、PR-AUC、召回率和F1更高。外部验证中，ACAWLR把误报从53个降到18个，精确率从0.595提高到0.788，但召回率从0.375下降到0.322。ACAWLR命中{ac_hits}/16个验证矿床，Logistic命中{log_hits}/16个；两者外部F1分别为0.457和0.460，几乎相同。")

    add_heading(doc, "7. 怎么解释最终结果", 1)
    add_body(doc, "如果目标是减少无效靶区、优先推荐少量高置信位置，ACAWLR更合适，因为它的排序指标和精确率更高、误报明显更少。如果目标是尽量不漏掉矿床，当前阈值下传统Logistic召回率和矿床命中数略高。")
    add_body(doc, "不能简单说“精确度提高了很多”。更准确的表述是：ACAWLR提高了排序能力和高置信预测质量，但在外部验证上以部分召回率为代价；两种模型的F1基本持平。")

    add_heading(doc, "8. 必须说明的限制", 1)
    add_body(doc, "正样本网格来自44个矿床，572个正网格并不等于572个独立矿床样本；背景点也不是真正的确认无矿点。IDW插值误差没有进入模型置信区间，采样稀疏地区不确定性更高。当前五折中每折验证集同时参与最佳epoch选择，因此五折指标可能轻微偏乐观；16点外部验证没有参与训练，参考价值更高。")
    add_body(doc, "点级传统GLM中，标准化log(Cu)的OR为4.53，95%CI为2.46–8.37，p=1.34×10⁻⁶，说明Cu与矿床标签存在强条件关联。但这不是因果证明，也不能解释为铜每增加1 ppm，矿床概率就增加4.53倍。")

    add_callout(
        doc,
        "建议用于汇报的结论",
        "在相同44/16矿床划分、相同网格样本和相同十变量下，ACAWLR较全局Logistic具有更好的矿床有利度排序能力，并显著减少外部验证误报；但传统Logistic在当前阈值下召回率和矿床命中数略高，两者F1基本相当。因此ACAWLR更适合强调高置信靶区的筛选，Logistic则保留更强的透明性和可解释性。",
    )

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    check = Document(path)
    empty_tables = sum(
        1 for table in check.tables
        if not any(cell.text.strip() for row in table.rows for cell in row.cells)
    )
    print(path)
    print(f"paragraphs={len(check.paragraphs)} tables={len(check.tables)} empty_tables={empty_tables}")
